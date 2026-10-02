"""Private in-memory report exports, shared by download routes and verification."""

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from flask import has_app_context

from app.services.reporting_service import local_datetime, schedule_totals, transaction_totals, wallet_change


MONEY_FORMAT = '"BDT "#,##0.00;[Red]-"BDT "#,##0.00'
DOCUMENTATION = (
    "Dates and times use Bangladesh time (UTC+6). Amounts and fees are in BDT.",
    "Amount is the transfer or payment principal. Fee is the separately charged service fee.",
    "Wallet change is the actual signed effect: incoming principal, or outgoing principal plus fee. Only SUCCESS records affect completed totals.",
    "Scheduled payments are instructions for future processing. Their amounts are reported separately and never added to completed transaction totals. COMPLETED schedules link to their ledger transaction.",
    "This report documents activity in the hackathon demo wallet. References identify individual records and support reconciliation.",
)


def _font_path():
    for path in (
        "C:/Windows/Fonts/Nirmala.ttc", "C:/Windows/Fonts/Nirmala.ttf", "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansBengali-Regular.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansBengaliUI-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ):
        if Path(path).is_file():
            return path
    return None


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    name = "WalletReport"
    if name not in pdfmetrics.getRegisteredFontNames():
        path = _font_path()
        if not path:
            return "Helvetica"
        pdfmetrics.registerFont(TTFont(name, path))
    return name


def _money(value):
    return f"BDT {Decimal(value or 0):,.2f}"


def _generated_at():
    return local_datetime(datetime.now(timezone.utc)).strftime("%d %b %Y, %H:%M") + " BD"


def _styles():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib import colors

    font = _pdf_font()
    styles = getSampleStyleSheet()
    for item in styles.byName.values():
        item.fontName = font
    styles.add(ParagraphStyle("ReportTitle", fontName=font, fontSize=21, leading=27, textColor=colors.HexColor("#17233c"), spaceAfter=12))
    styles.add(ParagraphStyle("ReportSection", fontName=font, fontSize=12, leading=17, textColor=colors.HexColor("#17233c"), spaceBefore=15, spaceAfter=8))
    styles.add(ParagraphStyle("ReportSmall", fontName=font, fontSize=8, leading=12, textColor=colors.HexColor("#4a566c"), spaceAfter=5, wordWrap="CJK"))
    styles.add(ParagraphStyle("ReportCell", fontName=font, fontSize=8, leading=11, wordWrap="CJK"))
    styles.add(ParagraphStyle("ReportHeader", fontName=font, fontSize=8, leading=11, textColor=colors.white, wordWrap="CJK"))
    for item in styles.byName.values():
        item.shaping = 1
    return styles


def _paragraph(value, style):
    from reportlab.platypus import Paragraph
    return Paragraph(escape(str(value or "-")).replace("\n", "<br/>"), style)


def _table(rows, widths, styles, *, header=True, row_padding=7):
    from reportlab.lib import colors
    from reportlab.platypus import LongTable, TableStyle

    prepared = [[_paragraph(value, styles["ReportHeader"] if header and i == 0 else styles["ReportCell"]) for value in row] for i, row in enumerate(rows)]
    result = LongTable(prepared, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    instructions = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), row_padding), ("BOTTOMPADDING", (0, 0), (-1, -1), row_padding),
        ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#dfe5ef")),
    ]
    if header:
        instructions.extend([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#233650")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fb")]),
        ])
    result.setStyle(TableStyle(instructions))
    return result


def _draw_pdf_brand(pdf, width, height):
    """Crisp vector version of the app's existing yellow/blue UpayX mark."""
    from reportlab.lib import colors

    pdf.saveState()
    pdf.translate(36, height - 64)
    pdf.scale(0.28, 0.28)
    pdf.setFillColor(colors.white)
    pdf.roundRect(0, 0, 120, 120, 28, fill=1, stroke=0)
    for x, color in ((40, "#ffd400"), (80, "#0b63ce")):
        pdf.setFillColor(colors.HexColor(color))
        pdf.circle(x, 88, 12, fill=1, stroke=0)
    pdf.setFillColor(colors.HexColor("#ffd400"))
    yellow = pdf.beginPath()
    yellow.moveTo(28, 72); yellow.lineTo(52, 72); yellow.lineTo(52, 48)
    yellow.curveTo(52, 37, 58, 31, 68, 31); yellow.lineTo(68, 13)
    yellow.curveTo(43, 13, 28, 26, 28, 48); yellow.close()
    pdf.drawPath(yellow, fill=1, stroke=0)
    pdf.setFillColor(colors.HexColor("#0b63ce"))
    blue = pdf.beginPath()
    blue.moveTo(68, 72); blue.lineTo(92, 72); blue.lineTo(92, 48)
    blue.curveTo(92, 26, 78, 13, 54, 13); blue.lineTo(54, 31)
    blue.curveTo(63, 31, 68, 36, 68, 48); blue.close()
    pdf.drawPath(blue, fill=1, stroke=0)
    pdf.restoreState()
    pdf.saveState()
    pdf.setFillColor(colors.HexColor("#0b63ce"))
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(78, height - 49, "UpayX")
    pdf.setStrokeColor(colors.HexColor("#dfe5ef"))
    pdf.line(36, height - 73, width - 36, height - 73)
    pdf.restoreState()


def _build_pdf(story, *, title, landscape=False):
    from reportlab.lib.pagesizes import A4, landscape as landscape_size
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas
    from reportlab.platypus import SimpleDocTemplate

    class NumberedCanvas(canvas.Canvas):
        """Add reliable page X of Y numbering after all pages are laid out."""
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.saved_pages = []

        def showPage(self):
            self.saved_pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            count = len(self.saved_pages)
            for state in self.saved_pages:
                self.__dict__.update(state)
                width, height = self._pagesize
                _draw_pdf_brand(self, width, height)
                self.setFont(_pdf_font(), 8)
                self.setFillColor(colors.HexColor("#586579"))
                self.drawString(36, 21, "Demo wallet report | Bangladesh time (UTC+6) | BDT")
                self.drawRightString(width - 36, 21, f"Page {self._pageNumber} of {count}")
                super().showPage()
            super().save()

    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape_size(A4) if landscape else A4,
        rightMargin=36, leftMargin=36, topMargin=88, bottomMargin=42,
        title=title, author="UpayX", allowSplitting=1,
    )
    document.build(story, canvasmaker=NumberedCanvas)
    return output.getvalue()


def receipt_fields(user, transaction):
    change = wallet_change(transaction)
    fields = [
        ("Account holder", user.full_name), ("Wallet number", user.mobile),
        ("Transaction ID", str(transaction.id)), ("Reference", transaction.reference or "-"),
        ("Date and time (BD)", local_datetime(transaction.created_at).strftime("%d %b %Y, %H:%M") + " BD"),
        ("Operation", transaction.title), ("Type", transaction.kind.replace("_", " ").title()),
        ("Counterparty / provider", transaction.counterparty or "Wallet"),
        ("Direction", "Money in" if transaction.direction == "IN" else "Money out"),
        ("Status", transaction.status.replace("_", " ").title()),
        ("Amount", _money(transaction.amount)), ("Fee", _money(transaction.fee)),
        ("Wallet change", _money(change)), ("Note", transaction.note or "No note"),
    ]
    if (has_app_context() and transaction.kind == "BILL_PAYMENT"
            and getattr(transaction, "id", None) and getattr(user, "id", None)
            and getattr(transaction, "user_id", None) == user.id):
        from app.domain.payment_plans import PaymentInvoice
        from app.services.service_catalog import BILL_CATEGORIES

        invoice = PaymentInvoice.query.filter_by(transaction_id=transaction.id, user_id=user.id).first()
        if invoice:
            fields[4:4] = [
                ("Invoice number", invoice.invoice_number),
                ("Bill account", invoice.account_reference),
                ("Bill invoice reference", invoice.invoice_reference or "Not supplied"),
                ("Category", BILL_CATEGORIES.get(invoice.category, {}).get("label", invoice.category)),
            ]
    return fields


def receipt_summary(transaction):
    if transaction.status != "SUCCESS":
        return f"This {transaction.status.lower()} record has no completed wallet deduction or credit."
    if transaction.direction == "IN":
        return f"{_money(transaction.amount)} was credited to the wallet."
    return f"{_money(transaction.amount)} principal and {_money(transaction.fee)} fee were deducted. Total wallet deduction: {_money(transaction.amount + (transaction.fee or 0))}."


def _receipt_pdf(title, fields, summary, status, reference):
    from reportlab.platypus import Spacer
    styles = _styles()
    story = [
        _paragraph(title, styles["ReportTitle"]),
        _paragraph(f"Generated: {_generated_at()}", styles["ReportSmall"]),
        _paragraph(status.replace('_', ' ').title(), styles["ReportSection"]),
        _table(fields, [170, 353], styles, header=False, row_padding=5),
        _paragraph("Record summary", styles["ReportSection"]),
        _paragraph(summary, styles["BodyText"]),
        Spacer(1, 8),
        _paragraph("Record documentation", styles["ReportSection"]),
    ]
    story.extend(_paragraph(text, styles["ReportSmall"]) for text in DOCUMENTATION)
    return _build_pdf(story, title=reference)


def _pdf_to_jpg(pdf_bytes):
    """Render the identical documented PDF, including shaped names, as a JPG."""
    import pymupdf
    from PIL import Image

    pages = []
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as document:
        for page in document:
            pixels = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False)
            pages.append(Image.frombytes("RGB", (pixels.width, pixels.height), pixels.samples))
    gap = 28
    output_image = Image.new("RGB", (max(page.width for page in pages), sum(page.height for page in pages) + gap * (len(pages) - 1)), "#e5eaf2")
    offset = 0
    for page in pages:
        output_image.paste(page, ((output_image.width - page.width) // 2, offset))
        offset += page.height + gap
    output = BytesIO()
    output_image.save(output, format="JPEG", quality=95, optimize=True, dpi=(180, 180))
    return output.getvalue()


def receipt_pdf(user, transaction, app_name):
    return _receipt_pdf(f"{app_name} transaction receipt", receipt_fields(user, transaction), receipt_summary(transaction), transaction.status, f"Transaction {transaction.reference or transaction.id}")


def receipt_jpg(user, transaction, app_name):
    return _pdf_to_jpg(receipt_pdf(user, transaction, app_name))


def schedule_receipt_fields(user, item):
    return [
        ("Account holder", user.full_name), ("Wallet number", user.mobile),
        ("Scheduled payment ID", str(item.id)),
        ("Due date and time (BD)", local_datetime(item.due_at).strftime("%d %b %Y, %H:%M") + " BD"),
        ("Type", item.kind.replace("_", " ").title()),
        ("Recipient / authorization", item.recipient_name or "-"),
        ("Recipient / account number", item.recipient_number),
        ("Provider", item.provider or "-"), ("Category", item.category or "-"),
        ("Amount", _money(item.amount)),
        ("Payment mode", "Auto pay" if item.auto_pay else "Manual payment"),
        ("Frequency", item.frequency.replace("_", " ").title()),
        ("Status", item.status.replace("_", " ").title()),
        ("Linked transaction ID", str(item.transaction_id or "No completed transaction")),
        ("Note", item.note or "No note"), ("Processing message", item.last_error or "-"),
    ]


def schedule_receipt_summary(item):
    if item.status == "COMPLETED":
        return f"The scheduled payment for {_money(item.amount)} was completed. Refer to linked transaction {item.transaction_id} for the actual wallet deduction and fees."
    return f"This {item.status.lower()} instruction records {_money(item.amount)} principal. This schedule record is separate from completed wallet transactions and does not itself prove that money was deducted."


def schedule_receipt_pdf(user, item, app_name):
    return _receipt_pdf(f"{app_name} scheduled payment", schedule_receipt_fields(user, item), schedule_receipt_summary(item), item.status, f"Scheduled payment {item.id}")


def schedule_receipt_jpg(user, item, app_name):
    return _pdf_to_jpg(schedule_receipt_pdf(user, item, app_name))


def _summary_rows(transactions, schedules):
    totals, planned = transaction_totals(transactions), schedule_totals(schedules)
    return [
        ("Transaction records", str(totals["transaction_count"])),
        ("Successful transactions", str(totals["completed_count"])),
        ("Other transaction statuses", str(totals["other_count"])),
        ("Completed money in", _money(totals["incoming_total"])),
        ("Completed money out including fees", _money(totals["outgoing_total"])),
        ("Completed fees charged", _money(totals["fees_total"])),
        ("Net completed wallet change", _money(totals["net_change"])),
        ("Scheduled payment records", str(planned["schedule_count"])),
        ("Awaiting scheduled payments", str(planned["scheduled_count"])),
        ("Planned principal awaiting payment", _money(planned["scheduled_total"])),
        ("Active auto-pay installments", str(planned["auto_pay_count"])),
    ]


def _summary_values(transactions, schedules):
    totals, planned = transaction_totals(transactions), schedule_totals(schedules)
    return [
        totals["transaction_count"], totals["completed_count"], totals["other_count"],
        totals["incoming_total"], totals["outgoing_total"], totals["fees_total"], totals["net_change"],
        planned["schedule_count"], planned["scheduled_count"], planned["scheduled_total"], planned["auto_pay_count"],
    ]


def _cache_formula_results(workbook_bytes, sheet_values):
    """Keep Excel formulas and their authoritative database snapshot values.

    openpyxl writes formulas without caches. Populate only the formulas created
    here, so previews/data-only readers show the same totals as the PDF before
    Excel's automatic recalculation refreshes them after an edit.
    """
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    output = BytesIO()
    with ZipFile(BytesIO(workbook_bytes)) as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            values = sheet_values.get(entry.filename)
            if values:
                tree = ElementTree.fromstring(content)
                for cell in tree.iter(f"{{{namespace}}}c"):
                    coordinate = cell.get("r")
                    if coordinate in values and cell.find(f"{{{namespace}}}f") is not None:
                        cached = cell.find(f"{{{namespace}}}v")
                        if cached is None:
                            cached = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                        cached.text = str(values[coordinate])
                        cell.set("t", "n")
                content = ElementTree.tostring(tree, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, content)
    return output.getvalue()


def report_pdf(user, transactions, schedules, filters, app_name):
    from reportlab.platypus import PageBreak
    styles = _styles()
    story = [
        _paragraph(f"{app_name} activity report", styles["ReportTitle"]),
        _paragraph(f"Account: {user.full_name} ({user.mobile})", styles["ReportSmall"]),
        _paragraph(f"Generated: {_generated_at()}", styles["ReportSmall"]),
        _paragraph(filters, styles["ReportSmall"]),
        _paragraph("Recorded transactions", styles["ReportSection"]),
    ]
    tx_rows = [["ID / Date (BD)", "Operation / Type", "Counterparty / Note", "Reference", "Direction / Status", "Amount (BDT)", "Fee (BDT)", "Wallet change (BDT)"]]
    for tx in transactions:
        tx_rows.append([
            f"{tx.id}\n{local_datetime(tx.created_at):%d %b %Y\n%H:%M}",
            f"{tx.title}\n{tx.kind}", f"{tx.counterparty or 'Wallet'}\n{tx.note or ''}", tx.reference or "-",
            f"{tx.direction}\n{tx.status}", f"{tx.amount:,.2f}", f"{tx.fee or 0:,.2f}", f"{wallet_change(tx):+,.2f}",
        ])
    if len(tx_rows) > 1:
        story.append(_table(tx_rows, [90, 105, 155, 108, 72, 75, 65, 100], styles))
    else:
        story.append(_paragraph("No transaction records match the selected filters.", styles["ReportSmall"]))
    story.append(_paragraph("Scheduled payments", styles["ReportSection"]))
    scheduled_rows = [["ID / Due date (BD)", "Type / Frequency", "Recipient / Provider", "Mode / Status", "Amount (BDT)", "Linked transaction / Notes"]]
    for item in schedules:
        scheduled_rows.append([
            f"{item.id}\n{local_datetime(item.due_at):%d %b %Y\n%H:%M}",
            f"{item.kind}\n{item.frequency}",
            f"{item.recipient_name or '-'}\n{item.recipient_number}\n{item.provider or ''}",
            f"{'Auto pay' if item.auto_pay else 'Manual payment'}\n{item.status}", f"{item.amount:,.2f}",
            f"{item.transaction_id or '-'}\n{item.note or ''}\n{item.last_error or ''}",
        ])
    if len(scheduled_rows) > 1:
        story.append(_table(scheduled_rows, [105, 120, 190, 110, 95, 150], styles))
    else:
        story.append(_paragraph("No scheduled payments match the selected filters.", styles["ReportSmall"]))
    story.append(_paragraph("Field documentation", styles["ReportSection"]))
    story.extend(_paragraph(text, styles["ReportSmall"]) for text in DOCUMENTATION)
    story.extend([
        PageBreak(), _paragraph("Report summary", styles["ReportTitle"]),
        _paragraph("Totals for the filtered records in this report", styles["ReportSmall"]),
        _table(_summary_rows(transactions, schedules), [400, 370], styles, header=False),
    ])
    story.append(_paragraph("Report filters: " + filters, styles["ReportSmall"]))
    return _build_pdf(story, title="Filtered wallet activity report", landscape=True)


def report_xlsx(user, transactions, schedules, filters, app_name):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties

    workbook = Workbook()
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    report = workbook.active
    report.title = "Transactions"
    report.sheet_properties.tabColor = "233650"
    planned = workbook.create_sheet("Scheduled payments")
    summary = workbook.create_sheet("Report summary")

    def text_cell(sheet, row, column, value):
        cell = sheet.cell(row=row, column=column, value=value)
        if isinstance(value, str):
            # Names, references and notes are untrusted text, including '=...'.
            cell.data_type = "s"
        return cell

    def formula_cell(sheet, row, column, formula, monetary=False):
        cell = sheet.cell(row=row, column=column, value=formula)
        cell.font = Font(name="Arial", size=10, color="17233C")
        cell.number_format = MONEY_FORMAT if monetary else "#,##0"
        cell.alignment = Alignment(horizontal="right", vertical="center")
        return cell

    def initialize(sheet, title, widths):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=.25, right=.25, top=.45, bottom=.45, header=.2, footer=.2)
        sheet.oddFooter.left.text = "Demo wallet report - BDT - Bangladesh time (UTC+6)"
        sheet.oddFooter.right.text = "Page &P of &N"
        sheet.oddFooter.left.size = sheet.oddFooter.right.size = 8
        sheet.oddHeader.right.text = "&A"
        for row, value in [(2, title), (3, f"Account: {user.full_name} ({user.mobile})"), (4, f"Generated: {_generated_at()}"), (5, filters)]:
            text_cell(sheet, row, 1, value)
            # Context is a presentation block and never merges data columns.
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(widths))
            sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = 28 if row == 5 else 21
        sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="17233C")
        for index, width in enumerate(widths, 1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(index)].width = width

    def add_rows(sheet, headers, rows, table_name):
        for column, header in enumerate(headers, 1):
            cell = text_cell(sheet, 7, column, header)
            cell.fill = PatternFill("solid", fgColor="233650")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[7].height = 32
        for row_index, values in enumerate(rows, 8):
            for column, value in enumerate(values, 1):
                cell = text_cell(sheet, row_index, column, value)
                cell.font = Font(name="Arial", size=10, color="17233C")
                cell.alignment = Alignment(horizontal="right" if isinstance(value, (int, float, Decimal)) else "left", vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="DFE5EF"))
                if isinstance(value, datetime):
                    cell.number_format = "dd mmm yyyy hh:mm"
                elif isinstance(value, Decimal):
                    cell.number_format = MONEY_FORMAT
            sheet.row_dimensions[row_index].height = 38 + max(0, max(len(str(value or "")) for value in values) // 45) * 12
        if rows:
            from openpyxl.utils import get_column_letter
            table = Table(displayName=table_name, ref=f"A7:{get_column_letter(len(headers))}{7 + len(rows)}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        else:
            text_cell(sheet, 8, 1, "No records match the selected filters.")
        sheet.freeze_panes = "C8"
        sheet.print_title_rows = "1:7"

    initialize(report, f"{app_name} recorded transactions", [8, 23, 23, 22, 30, 22, 11, 15, 18, 17, 20, 40])
    tx_rows = []
    for tx in transactions:
        tx_rows.append([
            tx.id, local_datetime(tx.created_at).replace(tzinfo=None), tx.title, tx.kind,
            tx.counterparty or "Wallet", tx.reference or "", tx.direction, tx.status,
            Decimal(tx.amount), Decimal(tx.fee or 0), wallet_change(tx), tx.note or "",
        ])
    add_rows(report, ["ID", "Date / time (BD)", "Operation", "Type", "Counterparty / provider", "Reference", "Direction", "Status", "Amount (BDT)", "Fee (BDT)", "Wallet change (BDT)", "Note"], tx_rows, "RecordedTransactions")
    for row in range(8, 8 + len(tx_rows)):
        formula_cell(report, row, 11, f'=IF(H{row}="SUCCESS",IF(G{row}="IN",I{row},-(I{row}+J{row})),0)', monetary=True)
    footer_row = 10 + len(tx_rows)
    report.row_breaks.append(Break(id=footer_row - 1))
    text_cell(report, footer_row, 1, "TRANSACTION SUMMARY")
    report.cell(footer_row, 1).font = Font(name="Arial", size=12, bold=True, color="17233C")
    for index, (label, _) in enumerate(_summary_rows(transactions, schedules)[:7], footer_row + 1):
        text_cell(report, index, 1, label)
        report.merge_cells(start_row=index, start_column=1, end_row=index, end_column=5)
        formula_cell(report, index, 6, f"='Report summary'!B{7 + index - footer_row - 1}", monetary=index >= footer_row + 4)
        report.row_dimensions[index].height = 22
    report.print_area = f"A1:L{footer_row + 7}"

    initialize(planned, f"{app_name} scheduled payments", [8, 23, 23, 24, 22, 25, 20, 17, 15, 18, 18, 40, 35])
    schedule_rows = []
    for item in schedules:
        schedule_rows.append([
            item.id, local_datetime(item.due_at).replace(tzinfo=None), item.kind,
            item.recipient_name or "", item.recipient_number, item.provider or "", item.category or "",
            item.frequency, "Auto pay" if item.auto_pay else "Manual", item.status,
            Decimal(item.amount), item.note or "", f"Transaction {item.transaction_id}" if item.transaction_id else item.last_error or "",
        ])
    add_rows(planned, ["ID", "Due date / time (BD)", "Type", "Recipient name", "Recipient / account number", "Provider", "Category", "Frequency", "Mode", "Status", "Amount (BDT)", "Note", "Linked transaction / error"], schedule_rows, "ScheduledPayments")
    planned.print_area = f"A1:M{max(8, 7 + len(schedule_rows))}"

    initialize(summary, f"{app_name} report summary", [48, 36, 32, 32])
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    tx_end = max(8, 7 + len(tx_rows))
    schedule_end = max(8, 7 + len(schedule_rows))
    tx_ids, tx_status, tx_direction = f"'Transactions'!A8:A{tx_end}", f"'Transactions'!H8:H{tx_end}", f"'Transactions'!G8:G{tx_end}"
    tx_amount, tx_fees, tx_change = f"'Transactions'!I8:I{tx_end}", f"'Transactions'!J8:J{tx_end}", f"'Transactions'!K8:K{tx_end}"
    schedule_ids, schedule_status = f"'Scheduled payments'!A8:A{schedule_end}", f"'Scheduled payments'!J8:J{schedule_end}"
    schedule_amount, schedule_mode = f"'Scheduled payments'!K8:K{schedule_end}", f"'Scheduled payments'!I8:I{schedule_end}"
    summary_formulas = [
        f"=COUNT({tx_ids})", f'=COUNTIFS({tx_status},"SUCCESS")', "=B7-B8",
        f'=SUMIFS({tx_amount},{tx_status},"SUCCESS",{tx_direction},"IN")',
        f'=SUMIFS({tx_amount},{tx_status},"SUCCESS",{tx_direction},"OUT")+SUMIFS({tx_fees},{tx_status},"SUCCESS",{tx_direction},"OUT")',
        f'=SUMIFS({tx_fees},{tx_status},"SUCCESS",{tx_direction},"OUT")',
        f"=SUM({tx_change})", f"=COUNT({schedule_ids})", f'=COUNTIFS({schedule_status},"SCHEDULED")',
        f'=SUMIFS({schedule_amount},{schedule_status},"SCHEDULED")',
        f'=COUNTIFS({schedule_status},"SCHEDULED",{schedule_mode},"Auto pay")',
    ]
    for row, ((label, _), formula) in enumerate(zip(_summary_rows(transactions, schedules), summary_formulas), 7):
        text_cell(summary, row, 1, label)
        formula_cell(summary, row, 2, formula, monetary=row in {10, 11, 12, 13, 16})
        summary.row_dimensions[row].height = 24
    documentation_row = 20
    text_cell(summary, documentation_row, 1, "FIELD DOCUMENTATION")
    summary.cell(documentation_row, 1).font = Font(name="Arial", size=12, bold=True, color="17233C")
    for row, note in enumerate(DOCUMENTATION, documentation_row + 1):
        text_cell(summary, row, 1, note)
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        summary.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row].height = 38
    summary.print_area = "A1:D25"
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.font == Font():
                    cell.font = Font(name="Arial", size=10, color="17233C")
    workbook.properties.title = "Filtered wallet activity report"
    workbook.properties.creator = app_name
    output = BytesIO()
    workbook.save(output)
    values = _summary_values(transactions, schedules)
    transaction_caches = {f"K{row}": wallet_change(tx) for row, tx in enumerate(transactions, 8)}
    transaction_caches.update({f"F{row}": value for row, value in enumerate(values[:7], footer_row + 1)})
    summary_caches = {f"B{row}": value for row, value in enumerate(values, 7)}
    return _cache_formula_results(output.getvalue(), {
        "xl/worksheets/sheet1.xml": transaction_caches,
        "xl/worksheets/sheet3.xml": summary_caches,
    })
