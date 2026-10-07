"""A judge's real-browser walkthrough of an isolated rich synthetic fixture.

Run: python -m tests.browser_submission_qa
The ordinary wallet and the reusable judge-demo.db are both preserved. Only a
new, independently created temporary judge fixture receives these UI actions.
"""

from io import BytesIO
from datetime import date
import hashlib
import json
import logging
from pathlib import Path
from threading import Thread

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from pypdf import PdfReader
from werkzeug.serving import make_server

from app.domain.ai_governance import AIConsent
from app.domain.models import Transaction, User
from app.domain.operations import ScheduledPayment
from app.domain.pilot import PilotParticipant
from app.extensions import db
from app.services.financial_health_service import health_for_user
from app.services.cashflow_model import forecast_for_user
from app.services.schedule_service import add_months
from scripts.judge_demo import ROOT, prepare_judge_demo
from scripts.provider_contract_demo import IsolatedDirectory


def run():
    logging.getLogger("werkzeug").setLevel(logging.ERROR)
    output = ROOT / "tmp" / "submission-qa" / "browser"
    output.mkdir(parents=True, exist_ok=True)
    preserved = {path: hashlib.sha256(path.read_bytes()).hexdigest()
                 for path in (ROOT / "instance" / "upay_hackathon.db", ROOT / "instance" / "judge-demo.db")
                 if path.exists()}
    directory = IsolatedDirectory("judge-submission-")
    app, server, worker = None, None, None
    checks, flows, downloads, errors, console_errors, network_requests = [], [], [], [], [], []
    report = {}
    origin = ""

    def layout(page, language, width, stage):
        page.evaluate("""async () => {
            await document.fonts.ready;
            document.documentElement.style.scrollBehavior = 'auto';
            document.activeElement?.blur();
            window.scrollTo({top:0,left:0,behavior:'instant'});
        }""")
        page.evaluate("() => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))")
        page.evaluate("() => window.scrollTo({top:0,left:0,behavior:'instant'})")
        metrics = page.evaluate("""() => ({width:innerWidth, scroll_y:scrollY,
            html_scroll:document.documentElement.scrollWidth,body_scroll:document.body.scrollWidth,
            language:document.documentElement.lang,
            overflow:Array.from(document.querySelectorAll('main,.page-heading,.section-card,.focus-card,.form-page-grid,.info-card'))
            .filter(e => {const r=e.getBoundingClientRect(); return r.right>innerWidth+1 || r.left < -1;})
            .map(e => e.className)})""")
        row = {"language": language, "viewport": width, "stage": stage,
               "path": page.url.replace(origin, ""), **metrics}
        checks.append(row)
        assert metrics["scroll_y"] == 0 and metrics["html_scroll"] <= width + 1 and metrics["body_scroll"] <= width + 1, row
        assert not metrics["overflow"] and metrics["language"] == language, row
        page.screenshot(path=str(output / f"{language}-{width}-{stage}.png"), full_page=True)

    def visit(page, path):
        response = page.goto(origin + path, wait_until="load")
        assert response.status == 200, (path, response.status)
        return response

    def on_console(message):
        if message.type != "error":
            return
        item = {"type": "console", "message": message.text}
        if any(term in message.text.lower() for term in ("content security policy", "refused to", "unsafe-inline", "violat")):
            errors.append(item)
        else:
            console_errors.append(item)  # Includes deliberately aborted recipient lookup.

    def pdf_download(page, selector, stem):
        with page.expect_download() as received:
            page.locator(selector).first.click()
        path = output / (stem + ".pdf")
        received.value.save_as(str(path))
        contents = path.read_bytes()
        assert contents.startswith(b"%PDF") and len(contents) > 1000
        parsed = PdfReader(BytesIO(contents))
        assert len(parsed.pages) >= 1
        text = "\n".join(p.extract_text() or "" for p in parsed.pages)
        assert "01329097775" in text
        downloads.append({"file": path.name, "bytes": len(contents), "pages": len(parsed.pages), "owner_verified": True})
        return text

    try:
        app, report = prepare_judge_demo(Path(directory.name) / "judge-review.db")
        assert report["created"] and report["dataset"]["days"] == 120 and report["model"]["forecast_available"]
        with app.app_context():
            owner = User.query.filter_by(mobile="01329097775").one()
            owner_id = owner.id
            initial_balance = owner.balance
            initial_transactions = Transaction.query.filter_by(user_id=owner_id).count()
            receipt_id = Transaction.query.filter_by(user_id=owner_id, status="SUCCESS").order_by(Transaction.created_at.desc()).first().id
            assert PilotParticipant.query.count() == 0 and AIConsent.query.count() == 0
        server = make_server("127.0.0.1", 0, app, threaded=True)
        worker = Thread(target=server.serve_forever, daemon=True)
        worker.start()
        origin = f"http://127.0.0.1:{server.server_port}"
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            for language in ("en", "bn"):
                for width in (1440, 390):
                    stage = []
                    context = browser.new_context(viewport={"width": width, "height": 960}, accept_downloads=True)
                    page = context.new_page()
                    page.on("pageerror", lambda exc: errors.append({"type": "pageerror", "message": str(exc)}))
                    page.on("console", on_console)
                    page.on("request", lambda request: network_requests.append(request.url))
                    visit(page, "/auth/login")
                    page.locator("[name=mobile]").fill("01329097775")
                    page.locator(".auth-card button[type=submit]").click()
                    assert "123456" in page.locator(".otp-demo").inner_text()
                    page.locator("[name=otp]").fill("123456")
                    page.locator(".auth-card button[type=submit]").click()
                    page.wait_for_url(origin + "/")
                    page.locator("#navbarLanguage").select_option(language)
                    page.wait_for_function("language => document.documentElement.lang === language", arg=language)
                    stage.append("explicit synthetic OTP sign-in")
                    visit(page, "/?days=30")
                    assert page.locator(".focus-workflows .focus-card").count() == 2
                    assert page.locator('.focus-workflows a[href="/insights"]').count() == 1
                    assert page.locator('.focus-workflows a[href="/schedules"]').count() == 1
                    assert page.locator(".transaction-list .transaction-row").count() > 0
                    layout(page, language, width, "two-task-dashboard")
                    stage.append("focused two-task dashboard with populated 30-day history")
                    visit(page, "/?days=1")
                    assert page.locator(".transaction-list .transaction-row").count() == 0
                    assert page.locator(".transaction-list .muted").inner_text().strip()
                    assert page.locator('.period-presets a[href="/?days=30"]').count() == 1
                    layout(page, language, width, "today-empty-with-date-controls")
                    stage.append("honest empty Today with access to earlier activity")
                    with app.app_context():
                        health = health_for_user(owner_id)
                        forecast = forecast_for_user(owner_id)
                    visit(page, "/insights")
                    assert f"{health['safe_to_spend']:,.2f}" in page.locator(".health-amount").inner_text()
                    assert f"{forecast['predicted_outflow']:,.2f}" in page.locator(".forecast-summary").inner_text()
                    assert page.locator(".health-commitment").count() == len(health["commitments"])
                    assert page.locator(".health-status-ok").count() == 1
                    assert page.locator("#planningReviewTitle").count() == 0  # No fabricated research consent.
                    if language == "bn":
                        assert "অটো পে পরিচালনা করুন" in page.locator('.section-head a[href="/schedules"]').inner_text()
                    layout(page, language, width, "weekly-plan-and-local-forecast")
                    stage.append("weekly commitments, calibrated local forecast and matched ledger")
                    page.locator(".forecast-card details > summary").click()
                    with page.expect_download() as evaluated:
                        page.locator('a[href="/insights/model-evaluation"]').click()
                    evaluation_path = output / f"{language}-{width}-model-evaluation.json"
                    evaluated.value.save_as(str(evaluation_path))
                    evaluation = json.loads(evaluation_path.read_text(encoding="utf-8"))
                    assert evaluation["model"]["mae_bdt"] > 0
                    downloads.append({"file": evaluation_path.name, "bytes": evaluation_path.stat().st_size, "synthetic_evaluation": True})
                    visit(page, "/schedules")
                    form = page.locator("form[data-schedule-form]")
                    if language == "bn":
                        assert "প্রতিটি কিস্তি রিপোর্টে দেখা যাবে।" in page.locator(".page-heading p").inner_text()
                        assert "পর্যন্ত প্রতি মাসে" in form.locator('[name=frequency] option[value="MONTHLY"]').inner_text()
                    form.locator("[name=recipient_number]").fill("DEMO-BLOCKED-1001")
                    page.wait_for_function("() => document.querySelector('[data-recipient-status]').dataset.state === 'blocked'")
                    assert form.locator("button[type=submit]").is_disabled()
                    assert form.locator("[name=recipient_number]").get_attribute("aria-invalid") == "true"
                    layout(page, language, width, "blocked-recipient-feedback")
                    stage.append("blocked recipient feedback prevents submission")
                    # Hold one actual request briefly to inspect the accessible loading state.
                    pending = []
                    route_pattern = "**/operations/recipient?**"
                    page.route(route_pattern, lambda route: pending.append(route))
                    form.locator("[name=recipient_number]").fill("DEMO-METER-1001")
                    page.wait_for_function("() => document.querySelector('[data-recipient-status]').dataset.state === 'loading'")
                    assert form.locator("[data-recipient-status]").get_attribute("aria-live") == "polite"
                    layout(page, language, width, "recipient-check-loading")
                    assert pending
                    with page.expect_response(lambda r: "/operations/recipient?" in r.url):
                        pending[-1].continue_()
                    page.unroute(route_pattern)
                    page.wait_for_function("() => document.querySelector('[data-recipient-status]').dataset.state === 'registered'")
                    stage.append("recipient lookup announces loading then verified account")
                    # A disconnected lookup has a useful retry/server-validation explanation.
                    page.route(route_pattern, lambda route: route.abort("failed"))
                    form.locator("[name=recipient_number]").fill("DEMO-METER-1002")
                    page.wait_for_function("() => document.querySelector('[data-recipient-status]').dataset.state === 'unavailable'")
                    assert form.locator("[data-recipient-status]").inner_text().strip()
                    layout(page, language, width, "recipient-check-unavailable")
                    page.unroute(route_pattern)
                    form.locator("[name=recipient_number]").fill("DEMO-METER-1001")
                    page.wait_for_function("() => document.querySelector('[data-recipient-status]').dataset.state === 'registered'")
                    stage.append("failed lookup gives useful connection/server-validation explanation")
                    form.locator("[name=amount]").fill("0")
                    assert not form.evaluate("form => form.checkValidity()")
                    form.locator("[name=amount]").fill("25")
                    form.locator("[name=frequency]").select_option("MONTHLY")
                    note = f"Synthetic judge walkthrough {language} {width}"
                    form.locator("[name=note]").fill(note)
                    preview = form.locator("[data-schedule-preview]").inner_text()
                    first_date = date.fromisoformat(form.locator("[name=due_date]").input_value())
                    last_date = date.fromisoformat(form.locator("[name=due_date]").get_attribute("max"))
                    expected_count = 1 + sum(add_months(first_date, offset) <= last_date for offset in (1, 2))
                    assert f"{25 * expected_count:,.2f}" in preview
                    assert (f"{expected_count} installment" if language == "en" else f"{expected_count} কিস্তি") in preview
                    assert form.evaluate("form => form.checkValidity()")
                    layout(page, language, width, "monthly-plan-preview")
                    with page.expect_navigation():
                        form.locator("button[type=submit]").click()
                    page.wait_for_url(origin + "/schedules")
                    with app.app_context():
                        plans = ScheduledPayment.query.filter_by(user_id=owner_id, note=note).order_by(ScheduledPayment.due_at).all()
                        assert len(plans) == expected_count and all(p.status == "SCHEDULED" and p.transaction_id is None for p in plans)
                        schedule_id = plans[0].id
                        assert db.session.get(User, owner_id).balance == initial_balance
                        assert Transaction.query.filter_by(user_id=owner_id).count() == initial_transactions
                    stage.append("monthly review and CSRF save create future plans without debit")
                    visit(page, f"/wallet/schedule/{schedule_id}/receipt")
                    assert note in page.locator(".receipt-card").inner_text()
                    layout(page, language, width, "saved-instruction-summary")
                    pdf_download(page, "form.report-export-form button[type=submit]", f"{language}-{width}-schedule-instruction")
                    stage.append("saved instruction summary and printable PDF")
                    visit(page, "/schedules")
                    with page.expect_navigation():
                        page.locator(f'form[action="/schedules/{schedule_id}/pay"] button[type=submit]').click()
                    with app.app_context():
                        assert db.session.get(ScheduledPayment, schedule_id).status == "SCHEDULED"
                        assert db.session.get(User, owner_id).balance == initial_balance
                    assert page.locator(".flash-stack").inner_text().strip()
                    stage.append("early payment request reports not due and leaves balance unchanged")
                    with page.expect_navigation():
                        page.locator(f'form[action="/schedules/{schedule_id}/cancel"] button[type=submit]').click()
                    with app.app_context():
                        assert db.session.get(ScheduledPayment, schedule_id).status == "CANCELLED"
                    stage.append("cancellation confirms no debit")
                    visit(page, f"/wallet/transaction/{receipt_id}")
                    layout(page, language, width, "existing-transaction-receipt")
                    pdf_download(page, "form.report-export-form button[type=submit]", f"{language}-{width}-wallet-receipt")
                    stage.append("existing owner receipt and printable PDF")
                    visit(page, "/wallet/history?days=30")
                    assert page.locator("#reportBarChart svg").count() == 1
                    layout(page, language, width, "filtered-report")
                    export_form = page.locator('form.report-export-form[action="/wallet/report/export"]').first
                    with page.expect_download() as exported:
                        export_form.locator("button[type=submit]").click()
                    workbook_path = output / f"{language}-{width}-wallet-report.xlsx"
                    exported.value.save_as(str(workbook_path))
                    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
                    assert workbook.sheetnames == ["Transactions", "Scheduled payments", "Report summary"]
                    assert workbook["Transactions"].max_row > 30
                    workbook.close()
                    downloads.append({"file": workbook_path.name, "bytes": workbook_path.stat().st_size, "sheets": 3})
                    export_form.locator("#reportFormat").select_option("pdf")
                    pdf_download(page, 'form.report-export-form[action="/wallet/report/export"] button[type=submit]', f"{language}-{width}-wallet-report")
                    stage.append("filtered chart report with Excel and PDF exports")
                    visit(page, "/wallet/history?q=JUDGE-NO-MATCHING-RECORDS&scope=transactions")
                    assert page.locator("#reportBarChart svg").count() == 0
                    assert page.locator("table tbody tr").filter(has=page.locator("td[colspan='9']")).count() == 1
                    layout(page, language, width, "report-empty-filters")
                    stage.append("empty filtered report offers reset and clear explanation")
                    visit(page, "/assistant/privacy")
                    assert page.locator('input[name=purpose][value="hosted_assistant"]').count() == 1
                    assert not page.locator('form[action="/assistant/privacy/consent"]').first.locator("[name=accept]").is_checked()
                    layout(page, language, width, "privacy-default-off")
                    stage.append("privacy purposes start separately disabled")
                    visit(page, "/assistant")
                    question = "How do I set up Auto Pay?" if language == "en" else "অটো পে কীভাবে চালু করব?"
                    page.locator(".assistant-page-card textarea[name=question]").fill(question)
                    with page.expect_navigation():
                        page.locator(".assistant-page-card form.form-stack button[type=submit]").click()
                    answer = page.locator(".assistant-page-messages .assistant-message-guide .assistant-answer-text").last.inner_text()
                    assert "Auto Pay" in answer or "অটো" in answer
                    mode = page.locator(".assistant-page-card .assistant-mode").inner_text()
                    assert "Local" in mode or "স্থানীয়" in mode
                    layout(page, language, width, "local-assistant-guidance")
                    stage.append("local bilingual guidance works without hosted consent or provider")
                    with app.app_context():
                        assert db.session.get(User, owner_id).balance == initial_balance
                        assert Transaction.query.filter_by(user_id=owner_id).count() == initial_transactions
                        assert PilotParticipant.query.count() == 0 and AIConsent.query.count() == 0
                    flows.append({"language": language, "viewport": width, "verified": stage,
                                  "wallet_unchanged": True, "real_customer_evidence_claimed": False})
                    context.close()
            browser.close()
        assert all(url.startswith(origin + "/") for url in network_requests), "Unexpected non-local network request"
        assert not errors, errors
    except Exception as exc:
        errors.append({"type": "walkthrough_failure", "exception": type(exc).__name__, "message": str(exc)})
        raise
    finally:
        if server is not None:
            server.shutdown()
            worker.join(timeout=5)
            server.server_close()
        if app is not None:
            with app.app_context():
                db.session.remove()
                db.engine.dispose()
        directory.cleanup()
        preserved_unchanged = all(hashlib.sha256(path.read_bytes()).hexdigest() == digest for path, digest in preserved.items())
        if not preserved_unchanged:
            errors.append({"type": "preserved_database_changed"})
        evidence = {"fixture": {"kind": "new_isolated_rich_synthetic_judge_fixture", "dataset": report.get("dataset"),
                        "ledger_reconciled": report.get("ledger_reconciled"), "model": report.get("model")},
                    "checks": checks, "flows": flows, "downloads": downloads, "errors": errors,
                    "console_errors": console_errors, "preserved_main_and_judge_files_unchanged": preserved_unchanged,
                    "csrf_enabled": True, "trusted_sessions_enabled": True, "hosted_ai_disabled": True,
                    "network_local_only": all(url.startswith(origin + "/") for url in network_requests)}
        (output / "results.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    assert not errors, errors
    print(f"Passed {len(checks)} judge walkthrough layout checks, {sum(len(flow['verified']) for flow in flows)} UI flows and {len(downloads)} file exports; preserved main and judge databases unchanged.", flush=True)


if __name__ == "__main__":
    run()
