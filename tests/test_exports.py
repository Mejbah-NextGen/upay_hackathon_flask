"""Regression coverage for private receipts, export filters and financial totals."""

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from unittest import TestCase

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from app.domain.models import Transaction, User
from app.domain.operations import ScheduledPayment
from app.extensions import db
from app.services.export_service import report_pdf
from app.services.reporting_service import LOCAL_TIMEZONE, filter_schedules, filter_transactions, local_datetime, schedule_totals, wallet_change
from tests.helpers import AppTestCase


class ExportRouteTests(AppTestCase):
    def setUp(self):
        super().setUp()
        self.login()
        ScheduledPayment.query.delete()
        self.now = datetime.now(timezone.utc).replace(microsecond=0)
        self.cash = self.add_tx("Cash receipt", kind="CASH_OUT", amount="100.00", fee="1.50")
        self.bill = self.add_tx("Gas receipt", kind="BILL_PAYMENT", amount="50.00", counterparty="Titas Gas", note="=HYPERLINK(\"https://example.invalid\")")
        self.failed = self.add_tx("Failed payment", kind="BILL_PAYMENT", amount="500.00", status="FAILED")
        self.old = self.add_tx("Old receipt", kind="MOBILE_RECHARGE", amount="25.00", created_at=self.now - timedelta(days=110))
        self.incoming = self.add_tx("Incoming receipt", kind="ADD_MONEY", amount="300.00", direction="IN")
        self.schedule = self.add_schedule("Future household payment", due_at=self.now + timedelta(days=20))
        self.completed = self.add_schedule("Completed schedule", status="COMPLETED", transaction_id=self.bill.id)
        other = User(full_name="Private Owner", mobile="01898000001", balance=Decimal("0.00"))
        db.session.add(other)
        db.session.flush()
        self.private_tx = self.add_tx("Private ledger detail", user_id=other.id)
        self.private_schedule = self.add_schedule("Private planned detail", user_id=other.id)
        db.session.commit()

    def add_tx(self, title, **values):
        fields = dict(user_id=self.user_id, kind="SEND_MONEY", direction="OUT", title=title, reference=f"REF-{title}", counterparty="Demo person", amount=Decimal("10.00"), fee=Decimal("0.00"), created_at=self.now, status="SUCCESS")
        fields.update(values)
        tx = Transaction(**fields)
        db.session.add(tx)
        db.session.flush()
        return tx

    def add_schedule(self, note, **values):
        fields = dict(user_id=self.user_id, kind="BILL_PAYMENT", recipient_number="DEMO-GAS-1001", recipient_name="Household account", provider="Titas Gas", category="gas", amount=Decimal("120.00"), note=note, frequency="ONE_TIME", auto_pay=True, due_at=self.now + timedelta(days=1), status="SCHEDULED", recurrence_group=note)
        fields.update(values)
        item = ScheduledPayment(**fields)
        db.session.add(item)
        db.session.flush()
        return item

    def test_receipt_access_is_owner_only_for_pages_and_both_formats(self):
        for suffix in ("", "/download?format=pdf", "/download?format=jpg"):
            self.assertEqual(404, self.client.get(f"/wallet/transaction/{self.private_tx.id}{suffix}").status_code)
        for suffix in ("/receipt", "/download?format=pdf", "/download?format=jpg"):
            self.assertEqual(404, self.client.get(f"/wallet/schedule/{self.private_schedule.id}{suffix}").status_code)
        self.assertEqual(404, self.client.get("/wallet/transaction/999999").status_code)
        with self.client.session_transaction() as state:
            state.clear()
        for url in ("/wallet/report", "/wallet/report/export?format=xlsx", f"/wallet/transaction/{self.cash.id}/download?format=jpg", f"/wallet/schedule/{self.schedule.id}/download?format=pdf"):
            response = self.client.get(url)
            self.assertEqual(302, response.status_code)
            self.assertIn("/auth/login", response.location)

    def test_transaction_receipts_document_actual_deduction_in_pdf_and_jpg(self):
        page = self.client.get(f"/wallet/transaction/{self.cash.id}")
        self.assertIn(b"101.50", page.data)
        response = self.client.get(f"/wallet/transaction/{self.cash.id}/download?format=pdf")
        self.assertEqual(200, response.status_code)
        self.assertEqual("application/pdf", response.mimetype)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertEqual("private, no-store", response.headers["Cache-Control"])
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(response.data)).pages)
        for expected in ("100.00", "1.50", "101.50", self.cash.reference, "Record documentation", "Page 1 of"):
            self.assertIn(expected, text)
        response = self.client.get(f"/wallet/transaction/{self.cash.id}/download?format=jpg")
        self.assertEqual("image/jpeg", response.mimetype)
        image = Image.open(BytesIO(response.data))
        image.verify()
        self.assertEqual("JPEG", image.format)
        self.assertGreater(image.height, image.width)

    def test_schedule_summary_documents_future_status_and_formats(self):
        response = self.client.get(f"/wallet/schedule/{self.schedule.id}/receipt")
        self.assertIn(b"Scheduled payment summary", response.data)
        self.assertIn(b"separate from completed wallet transactions", response.data)
        response = self.client.get(f"/wallet/schedule/{self.schedule.id}/download?format=pdf")
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(response.data)).pages)
        self.assertIn("Scheduled payment ID", text)
        self.assertIn("Scheduled", text)
        self.assertIn("DEMO-GAS-1001", text)
        response = self.client.get(f"/wallet/schedule/{self.schedule.id}/download?format=jpg")
        Image.open(BytesIO(response.data)).verify()

    def test_filtered_excel_keeps_typed_data_safe_text_and_paged_summary(self):
        response = self.client.get("/wallet/report/export?format=xlsx&kind=BILL_PAYMENT&status=SUCCESS&scope=transactions&days=7&q=titas")
        self.assertEqual(200, response.status_code)
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(["Transactions", "Scheduled payments", "Report summary"], book.sheetnames)
        sheet = book["Transactions"]
        self.assertEqual(self.bill.id, sheet["A8"].value)
        self.assertIsInstance(sheet["B8"].value, datetime)
        self.assertEqual(50, sheet["I8"].value)
        self.assertEqual("s", sheet["L8"].data_type)
        self.assertEqual(self.bill.note, sheet["L8"].value)
        self.assertEqual("f", sheet["K8"].data_type)
        self.assertNotIn("Old receipt", str(list(sheet.values)))
        self.assertNotIn("Private ledger detail", str(list(sheet.values)))
        self.assertEqual("landscape", sheet.page_setup.orientation)
        self.assertEqual("C8", sheet.freeze_panes)
        self.assertIn("Page &P of &N", sheet.oddFooter.right.text)
        self.assertEqual("f", book["Report summary"]["B11"].data_type)
        self.assertIn('"SUCCESS"', book["Report summary"]["B11"].value)
        self.assertIn("No records", book["Scheduled payments"]["A8"].value)
        cached = load_workbook(BytesIO(response.data), data_only=True)
        self.assertEqual(-50, cached["Transactions"]["K8"].value)
        self.assertEqual(50, cached["Report summary"]["B11"].value)
        self.assertEqual(1, cached["Report summary"]["B8"].value)
        self.assertEqual(0, cached["Report summary"]["B16"].value)

    def test_filtered_pdf_is_paginated_ends_with_summary_and_excludes_other_users(self):
        response = self.client.get("/wallet/report/export?format=pdf&days=7")
        reader = PdfReader(BytesIO(response.data))
        self.assertGreaterEqual(len(reader.pages), 2)
        text = " ".join(page.extract_text() for page in reader.pages)
        self.assertNotIn("Private ledger detail", text)
        self.assertNotIn("Private planned detail", text)
        self.assertNotIn("Old receipt", text)
        self.assertIn("Future household payment", text)
        final = reader.pages[-1].extract_text()
        self.assertIn("Report summary", final)
        self.assertIn("BDT 151.50", final)
        self.assertIn("BDT 300.00", final)
        self.assertIn("BDT 120.00", final)
        for number, page in enumerate(reader.pages, 1):
            self.assertIn(f"Page {number} of {len(reader.pages)}", page.extract_text())
            self.assertIn("UpayX", page.extract_text())

    def test_receipt_header_has_existing_vector_brand_in_reserved_space(self):
        import pymupdf
        response = self.client.get(f"/wallet/transaction/{self.cash.id}/download?format=pdf")
        with pymupdf.open(stream=response.data, filetype="pdf") as document:
            for page in document:
                wordmark = min(page.search_for("UpayX"), key=lambda rectangle: rectangle.y0)
                self.assertLess(wordmark.y1, 73)
                logo_shapes = [shape for shape in page.get_drawings() if shape["rect"].y1 < 73 and shape.get("fill")]
                self.assertGreaterEqual(len(logo_shapes), 5)

    def test_explicit_dates_status_and_schedule_scope_share_export_filters(self):
        today = local_datetime(self.now).date().isoformat()
        response = self.client.get(f"/wallet/report/export?format=xlsx&start_date={today}&end_date={today}&status=FAILED&scope=transactions")
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(self.failed.id, book["Transactions"]["A8"].value)
        response = self.client.get("/wallet/report/export?format=xlsx&scope=scheduled&status=SCHEDULED")
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(self.schedule.id, book["Scheduled payments"]["A8"].value)
        self.assertIn("No records", book["Transactions"]["A8"].value)
        self.assertNotIn("Private planned detail", str(list(book["Scheduled payments"].values)))
        cached = load_workbook(BytesIO(response.data), data_only=True)
        self.assertEqual(120, cached["Report summary"]["B16"].value)
        self.assertEqual(0, cached["Report summary"]["B11"].value)

    def test_invalid_formats_and_date_ranges_are_rejected(self):
        for url in ("/wallet/report/export?format=csv", f"/wallet/transaction/{self.cash.id}/download?format=exe", f"/wallet/schedule/{self.schedule.id}/download?format=png", "/wallet/report?start_date=2026-99-99", "/wallet/report/export?start_date=2026-10-02&end_date=2026-01-01"):
            self.assertEqual(400, self.client.get(url).status_code)

    def test_unicode_account_names_survive_pdf_and_jpg_downloads(self):
        self.user.full_name = "María Zoë রহমান"
        db.session.commit()
        response = self.client.get(f"/wallet/transaction/{self.cash.id}/download?format=pdf")
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(response.data)).pages)
        self.assertIn("María Zoë", text)
        self.assertIn("রহমান", text)
        response = self.client.get(f"/wallet/transaction/{self.cash.id}/download?format=jpg")
        image = Image.open(BytesIO(response.data))
        self.assertGreaterEqual(image.width, 1400)
        image.verify()

    def test_all_filtered_excel_cached_totals_match_completed_ledger_and_plans(self):
        response = self.client.get("/wallet/report/export?format=xlsx&days=7")
        cached = load_workbook(BytesIO(response.data), data_only=True)["Report summary"]
        self.assertEqual([4, 3, 1, 300, 151.5, 1.5, 148.5, 2, 1, 120, 1], [cached.cell(row, 2).value for row in range(7, 18)])

    def test_120_day_report_preserves_older_records_and_legacy_alias(self):
        for path in ("/wallet/report?days=120", "/wallet/history?days=120"):
            response = self.client.get(path)
            self.assertEqual(200, response.status_code)
            self.assertIn(b"Old receipt", response.data)
            self.assertIn(b"Print / download", response.data)
            self.assertNotIn(b"Private ledger detail", response.data)


class ExportCalculationTests(TestCase):
    def test_120_day_calendar_boundary_and_failed_wallet_effect(self):
        from types import SimpleNamespace
        now = datetime(2026, 10, 2, tzinfo=timezone.utc)
        item = SimpleNamespace(created_at=now - timedelta(days=119), kind="CASH_OUT", direction="OUT", status="SUCCESS", amount=Decimal("1.00"), fee=Decimal("0.02"), title="Test", counterparty="", reference="", note="")
        self.assertEqual([item], filter_transactions([item], days=120, now=now))
        self.assertEqual(Decimal("-1.02"), wallet_change(item))
        item.status = "FAILED"
        self.assertEqual(Decimal("0.00"), wallet_change(item))

    def test_planned_totals_do_not_include_completed_or_cancelled_installments(self):
        from types import SimpleNamespace
        due = datetime(2026, 11, 2, tzinfo=LOCAL_TIMEZONE)
        items = [SimpleNamespace(id=index, due_at=due, kind="BILL_PAYMENT", status=status, amount=Decimal("20.00"), auto_pay=True, recipient_name="Household", recipient_number="A1", provider="Gas", category="gas", note="") for index, status in enumerate(("SCHEDULED", "COMPLETED", "CANCELLED"), 1)]
        self.assertEqual(Decimal("20.00"), schedule_totals(items)["scheduled_total"])
        self.assertEqual([], filter_schedules(items, direction="IN"))
        self.assertEqual([items[0]], filter_schedules(items, query="HOUSEHOLD", status="SCHEDULED"))
        self.assertEqual([], filter_schedules(items, end_date=due.date() - timedelta(days=1)))
