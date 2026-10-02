import unittest
from datetime import datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import json
from types import SimpleNamespace
from urllib.parse import urlencode

from flask import template_rendered
from openpyxl import load_workbook
from pypdf import PdfReader

from app.domain.models import Transaction, User
from app.domain.operations import ScheduledPayment
from app.domain.payment_plans import PaymentInvoice
from app.extensions import db
from app.services.reporting_service import (
    LOCAL_TIMEZONE, dashboard_report, filter_chart_transactions, filter_transactions, local_datetime, parse_chart_selection, parse_days,
    period_dates, report_visualization,
)
from tests.helpers import AppTestCase


def transaction(created_at, *, kind="BILL_PAYMENT", amount="10.00", fee="0.00", direction="OUT", status="SUCCESS", **details):
    return SimpleNamespace(
        created_at=created_at, kind=kind, amount=Decimal(amount), fee=Decimal(fee),
        direction=direction, status=status, title=details.get("title", "Gas Bill Payment"),
        counterparty=details.get("counterparty", "Titas Gas"),
        reference=details.get("reference", "UPX-TEST"), note=details.get("note", ""),
    )


class ReportingUnitTests(unittest.TestCase):
    now = datetime(2026, 10, 2, 1, 0, tzinfo=timezone.utc)

    def test_today_uses_bangladesh_midnight_for_naive_and_aware_utc_storage(self):
        before = transaction(datetime(2026, 10, 1, 17, 59, 59))
        start = transaction(datetime(2026, 10, 1, 18, 0, 0))
        end = transaction(datetime(2026, 10, 2, 17, 59, 59, tzinfo=timezone.utc))
        tomorrow = transaction(datetime(2026, 10, 2, 18, 0, 0, tzinfo=timezone.utc))
        selected = filter_transactions([before, start, end, tomorrow], days=1, now=self.now)
        self.assertEqual(selected, [end, start])
        self.assertEqual(local_datetime(start.created_at).hour, 0)
        self.assertEqual(local_datetime(start.created_at).date(), self.now.date())

    def test_120_days_includes_first_calendar_day_and_excludes_day_before(self):
        start, end = period_dates(120, self.now)
        boundary = datetime.combine(start, time.min, LOCAL_TIMEZONE).astimezone(timezone.utc)
        first = transaction(boundary)
        outside = transaction(boundary - timedelta(microseconds=1))
        today = transaction(self.now)
        self.assertEqual(filter_transactions([first, outside, today], days=120, now=self.now), [today, first])
        self.assertEqual((end - start).days, 119)

    def test_day_range_clamps_and_invalid_values_have_safe_defaults(self):
        for raw, expected in [("0", 1), ("-3", 1), ("999", 120), ("12", 12), ("1.5", 1), ("invalid", 1), (None, 1)]:
            with self.subTest(raw=raw):
                self.assertEqual(parse_days(raw), expected)
        self.assertIsNone(parse_days("", default=None))

    def test_payment_totals_exclude_transfers_cash_out_and_incomplete_payments(self):
        rows = [
            transaction(self.now, amount="100.00"),
            transaction(self.now, kind="MOBILE_RECHARGE", amount="20.00"),
            transaction(self.now, kind="SEND_MONEY", amount="50.00"),
            transaction(self.now, kind="CASH_OUT", amount="100.00", fee="1.50"),
            transaction(self.now, kind="ADD_MONEY", amount="200.00", direction="IN"),
            transaction(self.now, amount="500.00", status="FAILED"),
            transaction(self.now, amount="400.00", status="PENDING"),
        ]
        report = dashboard_report(Decimal("5000.00"), rows, now=self.now)
        self.assertEqual(report["payment_total"], Decimal("120.00"))
        self.assertEqual(report["outgoing_total"], Decimal("271.50"))
        self.assertEqual(report["incoming_total"], Decimal("200.00"))
        self.assertEqual(report["transaction_count"], 7)
        self.assertEqual(report["balance"], Decimal("5000.00"))

    def test_every_dashboard_history_value_and_recent_rows_share_the_selected_period(self):
        rows = [transaction(self.now - timedelta(days=index), amount="10.00") for index in range(12)]
        one_day = dashboard_report("1234.56", rows, days=1, now=self.now)
        week = dashboard_report("1234.56", rows, days=7, now=self.now)
        all_rows = dashboard_report("1234.56", rows, days=90, now=self.now)
        self.assertEqual(one_day["transaction_count"], 1)
        self.assertEqual(week["transaction_count"], 7)
        self.assertEqual(week["payment_total"], Decimal("70.00"))
        self.assertEqual(all_rows["transaction_count"], 12)
        self.assertEqual(week["recent"], rows[:6])
        self.assertEqual(one_day["balance"], week["balance"])
        self.assertEqual(week["balance"], all_rows["balance"])

    def test_history_combines_day_direction_type_and_case_insensitive_search(self):
        match = transaction(self.now, counterparty="Titas Gas 12345", reference="UPX-BILL")
        recharge = transaction(self.now, kind="MOBILE_RECHARGE", title="Mobile Recharge", counterparty="Robi 01812345678")
        transfer = transaction(self.now, kind="SEND_MONEY", counterparty="Titas Gas 12345")
        old = transaction(self.now - timedelta(days=5), counterparty="Titas Gas 12345")
        rows = [match, recharge, transfer, old]
        self.assertEqual(filter_transactions(rows, days=1, direction="OUT", kind="payments", query="tItAs", now=self.now), [match])
        self.assertEqual(filter_transactions(rows, kind="MOBILE_RECHARGE", query="018123", now=self.now), [recharge])
        self.assertEqual(filter_transactions(rows, query="upx-bill", now=self.now), [match])

    def test_chart_series_share_successful_totals_and_bangladesh_calendar(self):
        rows = [
            transaction(datetime(2026, 10, 1, 18, 0), amount='100.00', fee='1.50'),
            transaction(self.now, amount='500.00', direction='IN', kind='ADD_MONEY'),
            transaction(self.now, amount='5000.00', status='FAILED'),
            transaction(self.now, amount='3000.00', status='PENDING'),
        ]
        charts = report_visualization(rows)
        self.assertEqual(charts['successful_count'], 2)
        self.assertEqual(charts['excluded_count'], 2)
        self.assertEqual(charts['daily'][0]['date'], '2026-10-02')
        self.assertEqual(charts['outgoing_total'], 101.5)
        self.assertEqual(charts['incoming_total'], 500.0)
        self.assertEqual(charts['daily'][-1]['cumulative_net'], 398.5)
        self.assertEqual(charts['categories'][0]['amount'], 101.5)
        self.assertEqual(sum(row['count'] for row in charts['histogram']), 1)
        self.assertEqual(charts['histogram'][-1]['cumulative_percent'], 100)

    def test_histogram_includes_every_boundary_amount_once_and_has_overflow_bin(self):
        rows = [transaction(self.now, amount=str(amount)) for amount in (499, 500, 999, 1000, 5000, 10000, 50000, 100000)]
        histogram = report_visualization(rows)['histogram']
        self.assertEqual([row['count'] for row in histogram], [1, 2, 1, 1, 1, 2])
        self.assertEqual([row['cumulative_count'] for row in histogram], [1, 3, 4, 5, 6, 8])
        self.assertEqual(histogram[-1]['cumulative_percent'], 100)
        self.assertEqual(report_visualization([])['daily'], [])
        self.assertEqual(report_visualization([])['histogram'][-1]['cumulative_percent'], 0)

    def test_chart_selections_use_local_dates_successful_ledger_and_fee_inclusive_bands(self):
        rows = [
            transaction(datetime(2026, 10, 1, 18), amount="499.00", fee="1.00"),
            transaction(datetime(2026, 10, 1, 17, 59), amount="500.00"),
            transaction(self.now, amount="500.00", status="FAILED"),
            transaction(self.now, amount="500.00", direction="IN", kind="ADD_MONEY"),
        ]
        selection = parse_chart_selection(json.dumps({"mode": "daily", "start": "2026-10-02", "end": "2026-10-02", "direction": "OUT"}))
        self.assertEqual(filter_chart_transactions(rows, selection), [rows[0]])
        selection = parse_chart_selection(json.dumps({"mode": "amount", "lower": 500, "upper": 1000}))
        self.assertEqual(filter_chart_transactions(rows, selection), rows[:2])
        selection = parse_chart_selection(json.dumps({"mode": "cumulative", "end": "2026-10-02", "direction": "IN"}))
        self.assertEqual(filter_chart_transactions(rows, selection), [rows[3]])

    def test_invalid_or_nonfinite_chart_selections_are_rejected(self):
        invalid = ["[]", "null", "not-json", json.dumps({"mode": "daily", "start": "2026-10-03", "end": "2026-10-02", "direction": "OUT"}), json.dumps({"mode": "amount", "lower": "NaN", "upper": 100}), json.dumps({"mode": "amount", "lower": 500, "upper": 500}), json.dumps({"mode": "category", "kind": "BILL_PAYMENT", "category": "unknown"}), json.dumps({"mode": "cumulative", "end": "invalid"})]
        for raw in invalid:
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                parse_chart_selection(raw)


class ReportingRouteTests(AppTestCase):
    def setUp(self):
        super().setUp()
        self.login()
        today = local_datetime(datetime.now(timezone.utc)).date()
        self.today_utc = datetime.combine(today, time(12), LOCAL_TIMEZONE).astimezone(timezone.utc).replace(tzinfo=None)
        self.add_tx("BILL_PAYMENT", "Today's Gas Bill", "100.00", self.today_utc, counterparty="Titas Gas")
        self.add_tx("MOBILE_RECHARGE", "Older Recharge", "25.00", self.today_utc - timedelta(days=2))
        self.add_tx("SEND_MONEY", "Old Transfer", "50.00", self.today_utc - timedelta(days=10))
        self.add_tx("CASH_OUT", "Today's Cash Out", "100.00", self.today_utc, fee="1.50")
        self.add_tx("BILL_PAYMENT", "Failed Bill", "500.00", self.today_utc, status="FAILED")
        other = User(full_name="Other User", mobile="01812345678", balance=Decimal("0.00"))
        db.session.add(other)
        db.session.flush()
        self.add_tx("BILL_PAYMENT", "Private Payment", "5000.00", self.today_utc, user_id=other.id)
        db.session.commit()

    def add_tx(self, kind, title, amount, created_at, **extra):
        tx = Transaction(
            user_id=extra.pop("user_id", self.user_id), kind=kind, direction="OUT", title=title,
            amount=Decimal(amount), created_at=created_at, reference=f"TEST-{title}", **extra,
        )
        db.session.add(tx)
        return tx

    def get_context(self, url):
        captured = []
        def record(sender, template, context, **extra):
            captured.append(context)
        template_rendered.connect(record, self.app)
        try:
            response = self.client.get(url)
        finally:
            template_rendered.disconnect(record, self.app)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(captured)
        return response, captured[-1]

    def test_dashboard_filter_and_history_links_keep_current_balance_and_scope(self):
        response, context = self.get_context("/?days=7")
        stats = context["stats"]
        self.assertEqual(stats["transaction_count"], 4)
        self.assertEqual(stats["payment_total"], Decimal("125.00"))
        self.assertEqual(stats["outgoing_total"], Decimal("226.50"))
        self.assertEqual(stats["balance"], Decimal("1000.00"))
        self.assertIn(b"days=7", response.data)
        page_content = response.data.split(b'<section class="page-content">', 1)[1]
        self.assertNotIn(b"Old Transfer", page_content)
        self.assertNotIn(b"Private Payment", response.data)
        self.assertIn(b"category=gas", response.data)
        _, today_context = self.get_context("/?days=1")
        self.assertEqual(today_context["stats"]["transaction_count"], 3)
        self.assertEqual(today_context["stats"]["payment_total"], Decimal("100.00"))

    def test_history_shows_filtered_records_fees_and_local_timestamp(self):
        response, context = self.get_context("/wallet/history?days=7&direction=OUT&kind=payments&q=titas")
        self.assertEqual(len(context["transactions"]), 1)
        self.assertEqual(context["totals"]["outgoing_total"], Decimal("100.00"))
        self.assertIn(b"12:00 PM", response.data)
        self.assertNotIn(b"Private Payment", response.data)
        _, cash_context = self.get_context("/wallet/history?days=1&kind=CASH_OUT")
        self.assertEqual(cash_context["totals"]["outgoing_total"], Decimal("101.50"))

    def test_bad_date_range_is_clamped_and_all_history_can_be_reset(self):
        _, context = self.get_context("/?days=500")
        self.assertEqual(context["days"], 120)
        _, context = self.get_context("/wallet/history")
        self.assertIsNone(context["days"])
        self.assertEqual(len(context["transactions"]), 5)
        _, context = self.get_context("/wallet/history?days=12")
        self.assertEqual(context["days"], 12)

    def test_filtered_charts_use_same_user_scope_and_records_as_report(self):
        response, context = self.get_context('/wallet/history?days=7&kind=CASH_OUT&direction=OUT')
        charts = context['analytics']
        self.assertEqual(charts['outgoing_total'], 101.5)
        self.assertEqual(charts['successful_count'], 1)
        self.assertEqual([row['kind'] for row in charts['categories']], ['CASH_OUT'])
        self.assertNotIn('Private Payment', response.get_data(as_text=True))
        self.assertIn('reportCumulativeChart', response.get_data(as_text=True))
        self.assertIn('reportHistogramChart', response.get_data(as_text=True))
        _, context = self.get_context('/wallet/history?status=FAILED')
        self.assertEqual(context['analytics']['successful_count'], 0)
        self.assertEqual(context['analytics']['outgoing_total'], 0)
        response, context = self.get_context('/wallet/history?scope=scheduled')
        self.assertNotIn('reportChartData', response.get_data(as_text=True))
        self.assertEqual(context['analytics']['successful_count'], 0)

    def test_pie_categories_use_owned_invoice_metadata_and_legacy_education_titles(self):
        school = self.add_tx("BILL_PAYMENT", "Shared provider fee", "60.00", self.today_utc, counterparty="Shared provider")
        education = self.add_tx("BILL_PAYMENT", "Education Payment", "70.00", self.today_utc, counterparty="Demo Examination Fees")
        db.session.flush()
        db.session.add(PaymentInvoice(user_id=self.user_id, transaction_id=school.id, invoice_number="SCHOOL-CHART", category="education-school", provider="Shared provider", account_reference="STUDENT-1"))
        db.session.commit()
        _, context = self.get_context("/wallet/report?kind=BILL_PAYMENT")
        categories = {row["category"]: row for row in context["analytics"]["categories"]}
        self.assertEqual(categories["education-school"]["label"], "School")
        self.assertEqual(categories["education"]["amount"], 70.0)
        _, context = self.get_context("/wallet/report?category=education-school")
        self.assertEqual([tx.id for tx in context["transactions"]], [school.id])
        _, context = self.get_context("/wallet/report?category=education")
        self.assertEqual({tx.id for tx in context["transactions"]}, {school.id, education.id})
        selection = {"mode": "category", "kind": "BILL_PAYMENT", "category": "education-school"}
        response = self.client.get("/wallet/report/export?" + urlencode({"format": "xlsx", "chart_selection": json.dumps(selection)}))
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        ids = [row[0] for row in workbook["Transactions"].iter_rows(min_row=8, values_only=True) if isinstance(row[0], int)]
        self.assertEqual(ids, [school.id])

    def test_auto_pay_reports_and_exports_exclude_manual_plans_and_unrelated_ledger(self):
        ScheduledPayment.query.delete()
        completed_tx = self.add_tx("BILL_PAYMENT", "Auto Plan Payment", "60.00", self.today_utc)
        manual_tx = self.add_tx("BILL_PAYMENT", "Manual Plan Payment", "70.00", self.today_utc)
        db.session.flush()
        plans = []
        for note, auto_pay, tx in [("Auto completed", True, completed_tx), ("Auto upcoming", True, None), ("Manual completed", False, manual_tx)]:
            item = ScheduledPayment(user_id=self.user_id, kind="BILL_PAYMENT", recipient_number="ACCOUNT-1", provider="Titas Gas", category="gas", amount=Decimal("60.00"), frequency="ONE_TIME", auto_pay=auto_pay, due_at=self.today_utc + timedelta(days=3), status="COMPLETED" if tx else "SCHEDULED", transaction_id=tx.id if tx else None, recurrence_group=note, note=note)
            db.session.add(item); plans.append(item)
        db.session.commit()
        response, context = self.get_context("/wallet/report?plan_mode=auto_pay")
        self.assertEqual([tx.id for tx in context["transactions"]], [completed_tx.id])
        self.assertEqual({item.id for item in context["schedules"]}, {item.id for item in plans[:2]})
        self.assertIn("Auto Pay Report", response.get_data(as_text=True))
        self.assertIn('name="plan_mode" value="auto_pay"', response.get_data(as_text=True))
        self.assertIn("plan_mode=auto_pay", self.client.get("/schedules").get_data(as_text=True))
        export = self.client.get("/wallet/report/export?format=xlsx&plan_mode=auto_pay")
        workbook = load_workbook(BytesIO(export.data), data_only=True)
        tx_ids = [row[0] for row in workbook["Transactions"].iter_rows(min_row=8, values_only=True) if isinstance(row[0], int)]
        plan_ids = [row[0] for row in workbook["Scheduled payments"].iter_rows(min_row=8, values_only=True) if isinstance(row[0], int)]
        self.assertEqual(tx_ids, [completed_tx.id])
        self.assertEqual(set(plan_ids), {item.id for item in plans[:2]})

    def test_segment_pdf_preserves_base_filters_and_excludes_failed_private_and_planned_records(self):
        selection = {"mode": "daily", "start": local_datetime(self.today_utc).date().isoformat(), "end": local_datetime(self.today_utc).date().isoformat(), "direction": "OUT"}
        params = {"format": "pdf", "days": 7, "kind": "CASH_OUT", "chart_selection": json.dumps(selection)}
        response = self.client.get("/wallet/report/export?" + urlencode(params))
        self.assertEqual(response.status_code, 200)
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(response.data)).pages)
        self.assertIn("Today's Cash Out", text)
        self.assertIn("101.50", text)
        for omitted in ("Today's Gas Bill", "Failed Bill", "Private Payment", "Older Recharge"):
            self.assertNotIn(omitted, text)
        self.assertEqual(self.client.get("/wallet/report/export?format=pdf&chart_selection=[]").status_code, 400)


if __name__ == "__main__":
    unittest.main()
