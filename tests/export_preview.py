"""Generate and render isolated export fixtures for visual QA.

Usage: python -m tests.export_preview
Requires requirements-dev.txt. Outputs go to tmp/export-qa.
"""

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pymupdf
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from app import create_app
from app.domain.models import User
from app.services.export_service import (
    receipt_jpg, receipt_pdf, schedule_receipt_jpg, schedule_receipt_pdf,
)
from tests.preview import PreviewConfig


def create_examples(output):
    # The normal fixtures are seeded into a fresh in-memory database.
    app = create_app(PreviewConfig)
    with app.app_context():
        user = User.query.filter_by(mobile="01329097775").one()
        client = app.test_client()
        with client.session_transaction() as state:
            state["user_id"] = user.id
        for extension in ("pdf", "xlsx"):
            response = client.get(f"/wallet/report/export?format={extension}&days=120")
            assert response.status_code == 200
            (output / f"report.{extension}").write_bytes(response.data)

        now = datetime.now(timezone.utc)
        transaction = SimpleNamespace(
            id=99999, reference="QA-LONG-REFERENCE-" + "X" * 50, title="Cash Out",
            created_at=now, status="SUCCESS", kind="CASH_OUT", direction="OUT",
            counterparty="Authorized recipient with a long account name " * 2,
            amount=Decimal("12345.67"), fee=Decimal("185.19"),
            note="This receipt documents a long note for wrapping and fee review. " * 4,
        )
        account = SimpleNamespace(full_name="Account holder with a long Unicode name María Zoë রহমান", mobile=user.mobile)
        for extension, generator in (("pdf", receipt_pdf), ("jpg", receipt_jpg)):
            (output / f"receipt.{extension}").write_bytes(generator(account, transaction, app.config["APP_NAME"]))
        schedule = SimpleNamespace(
            id=99999, kind="BILL_PAYMENT", recipient_number="ACCOUNT-12345678901234567890123456789",
            recipient_name="Household authorization owner with a long name " * 2, provider="DESCO Electricity",
            category="electricity", amount=Decimal("1200.00"), due_at=now + timedelta(days=30),
            note="Upcoming scheduled payment for household expenses. " * 4,
            auto_pay=True, status="SCHEDULED", frequency="MONTHLY", transaction_id=None, last_error=None,
        )
        for extension, generator in (("pdf", schedule_receipt_pdf), ("jpg", schedule_receipt_jpg)):
            (output / f"schedule.{extension}").write_bytes(generator(account, schedule, app.config["APP_NAME"]))


def render_examples(output):
    for name in ("report", "receipt", "schedule"):
        thumbnails = []
        with pymupdf.open(output / f"{name}.pdf") as document:
            for index, page in enumerate(document):
                page_path = output / f"{name}-page-{index + 1}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(1.6, 1.6)).save(page_path)
                for word in page.get_text("words"):
                    assert word[0] >= 0 and word[1] >= 0 and word[2] <= page.rect.width + 1 and word[3] <= page.rect.height + 1, (name, index, word)
                with Image.open(page_path) as source:
                    thumbnail = source.copy()
                thumbnail.thumbnail((630, 520))
                tile = Image.new("RGB", (650, 555), "#d9dee6")
                tile.paste(thumbnail, ((650 - thumbnail.width) // 2, 22))
                ImageDraw.Draw(tile).text((16, 533), f"{name} page {index + 1} of {len(document)}", fill="black")
                thumbnails.append(tile)
            print(name, len(document), "PDF pages, text inside bounds")
        contact = Image.new("RGB", (1300, 555 * ((len(thumbnails) + 1) // 2)), "white")
        for index, tile in enumerate(thumbnails):
            contact.paste(tile, (650 * (index % 2), 555 * (index // 2)))
        contact.save(output / f"{name}-contact.png")


def main():
    output = Path("tmp/export-qa")
    output.mkdir(parents=True, exist_ok=True)
    create_examples(output)
    render_examples(output)
    cached = load_workbook(output / "report.xlsx", data_only=True)
    print("Cached summary", [cached["Report summary"].cell(row, 2).value for row in range(7, 18)])


if __name__ == "__main__":
    main()
